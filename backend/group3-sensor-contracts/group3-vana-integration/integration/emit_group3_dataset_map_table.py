# -*- coding: utf-8 -*-
"""Emit Group 3's row-by-row status against Sir's DATASET CONNECTION MAP FOR CLOSING.

WHY THIS EXISTS
---------------
Sir, 2026-08-29: "Also has group 3 tabulated the message shared on Hackaverse group"
and earlier: "make a table of this marked message. THIS WILL GIVE YOU EXACT DETAILS
OF WHAT MUST BE DONE TO CLOSE THIS."

Format follows the precedent already set in the group: Dataset | Status |
Can you call it DONE? | Why.

SCOPE: Group 3's own rows only. Group 3 does not report status for another
group's rows - that is theirs to report. Group 3 appears in three of Sir's G1
rows as the SOURCE, so those are included and marked as such.

Every status is drawn from evidence on disk in this workspace.

Run:  python emit_group3_dataset_map_table.py
"""
import datetime
import io
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DAY6 = os.path.dirname(HERE)

# ---------------------------------------------------------------- Sir's rows
# (group, dataset, take_it_from, where_it_goes, status, done, why)
ROWS = [
    # --- Sir's G3 table ---
    ("G3", "Real available sensor measurements",
     "Hemanth + physically available hardware", "VANA intake -> G1",
     "RED  Not done", "NO",
     "Zero physical sensor observations exist. Camera confirmed present by Suraj 2026-08-29; "
     "defect G-42 destroys the first captured frame while reporting success (reproduced "
     "2304 bytes -> 0). Three-line fix written and regression-tested, with Dhruv. "
     "Group 3 will not emit an observation whose artifact is empty."),

    ("G3", "SAMACHAR data",
     "Omkar Patil + Samachar team", "VANA intake -> G1",
     "RED  Waiting on the source", "NO",
     "A SAMACHAR observation artifact is reported to exist but has not reached Group 3 intake. "
     "The intake path and endpoint already exist and are proven; no build is required. "
     "Group 3 has requested the artifact. Turnaround once received is minutes."),

    ("G3", "Marine Task 3 data",
     "Shravani / Marine team", "Observation/context pipeline AS APPLICABLE",
     "GREY  Qualified - awaiting one answer", "NOT APPLICABLE UNLESS OBSERVATION-SHAPED",
     "Sir's wording for Group 3 is 'as applicable'. Its firm destination in the map is the "
     "context layer, not Group 1 ingestion. Group 3 has asked one question: do any records "
     "carry a measurement + unit + location + timestamp? Only those belong in Group 3's path."),

    ("G3", "National GIS data",
     "Shravani / GIS team", "Geographic/context ENRICHMENT",
     "GREY  Qualified - awaiting one answer", "NOT APPLICABLE UNLESS OBSERVATION-SHAPED",
     "Sir's wording for Group 3 is 'enrichment'. Same question asked as for Marine Task 3. "
     "If the content is boundaries, layers or spatial truth, nothing is required from "
     "Group 3 and the row closes as not-applicable."),

    ("G3", "GOUDHA data",
     "GOUDHA (a system/product - owner not yet identified)", "Supporting data WHERE APPLICABLE",
     "GREY  Qualified - owner unidentified", "NOT APPLICABLE UNLESS OBSERVATION-SHAPED",
     "Sir's wording for Group 3 is 'where applicable'. Group 3 has not identified who holds "
     "the 41 records, so the request cannot yet be addressed to anyone. Same "
     "observation-shaped test applies."),

    # --- Sir's G1 rows where GROUP 3 IS THE SOURCE ---
    ("G1 (Group 3 is the source)", "Thane Creek Open-Meteo observation",
     "Existing G1 runtime", "MasterDB / canonical layer",
     "GREEN  Live and verified", "YES",
     "TC-Z03-EXT-OPENMETEO-OBS001 -> canonical record CR-b4615a27-7ab1-4bde-a078-a56fa0f2414c. "
     "Emitted by Group 3 from a real external source, 13/13 live acceptance 2026-08-25T11:04:54Z, "
     "re-verified by GET since. Identity, both timestamps, coordinates and provenance all "
     "preserved into the canonical record."),

    ("G1 (Group 3 is the source)", "Real observations currently available",
     "Group 3 / Hemanth + sensor/data owners", "G1 canonical ingestion -> MasterDB",
     "GREEN  Delivered for the data path", "YES for data source / NO for physical sensor",
     "One real external data source is live end to end. No physical sensor observation exists - "
     "see the sensor row above. Sir's acceptance criterion names 'physical/data source'; the "
     "data-source half is met."),

    ("G1 (Group 3 is the source)", "Mumbai, Navi Mumbai, Vasai, Thane, Maval observations",
     "Group 3 / regional owners", "Must be genuinely ingested, not recreated",
     "RED  Blocked at ingestion", "NO",
     "All five are built, schema- and semantically valid, and byte-preserved (sha256 verified "
     "against the stored source bytes, 5 of 5). POST returns HTTP 500 for a previously-unseen "
     "survey - three attempts, latest 2026-08-29T09:58:39Z. Payload proven not to be the cause: "
     "identical key set, zero type mismatches, 25 of 37 fields byte-identical against the record "
     "that passed Group 1's own 13/13 acceptance. Nothing was created. Group 3 has not retried."),
]

SINGLE_JOB = ("Make real external/sensor data consumable through the VANA intake contract.")


def verify():
    """Re-check the load-bearing numbers rather than trusting the text above."""
    out = {}
    sch = json.load(io.open(os.path.join(DAY6, "contract", "observation.schema.v2.2.json"),
                            encoding="utf-8"))
    out["required_fields"] = len(sch["required"])
    out["cross_field_rules"] = len(sch["allOf"])

    import glob, hashlib
    ok = 0
    for r in ["MU", "NM", "VS", "THN", "MV"]:
        p = os.path.join(DAY6, "evidence", "live", "live_mission_package_%s.json" % r)
        o = json.load(io.open(p, encoding="utf-8"))["observations"][0]
        h = o["raw_artifact_integrity"]["checksum_sha256"]
        c = glob.glob(os.path.join(DAY6, "evidence", "live", "raw_artifact_%s*.json" % h[:16]))
        if c and hashlib.sha256(io.open(c[0], "rb").read()).hexdigest() == h:
            ok += 1
    out["regions_byte_verified"] = ok

    phys = 0
    for f in glob.glob(os.path.join(DAY6, "evidence", "**", "*.json"), recursive=True):
        if '"hardware_verified": true' in io.open(f, encoding="utf-8", errors="replace").read():
            phys += 1
    out["physical_observations"] = phys
    return out


def main():
    v = verify()
    if v["regions_byte_verified"] != 5:
        raise SystemExit("ABORT: expected 5 byte-verified regions, got %d"
                         % v["regions_byte_verified"])
    if v["physical_observations"] != 0:
        raise SystemExit("ABORT: expected 0 physical observations, found %d"
                         % v["physical_observations"])

    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        raise SystemExit("ABORT: openpyxl not installed.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "G3 Dataset Map"

    title = "GROUP 3 - status against VANA DATASET CONNECTION MAP FOR CLOSING"
    ws.append([title])
    ws.append(["Compiled by Hemanth Reddy, Group 3 Integration - %s"
               % datetime.date.today().isoformat()])
    ws.append(["Group 3's own rows only. Status for another group's rows is theirs to report."])
    ws.append([])
    header = ["Group", "Dataset / data source", "Take it from", "Where it goes",
              "Group 3 status", "Can you call it DONE?", "Why"]
    ws.append(header)

    for r in ROWS:
        ws.append(list(r))

    ws.append([])
    ws.append(["G3's single job (Sir's words)", SINGLE_JOB])
    ws.append([])
    ws.append(["Verified at generation:",
               "V2.2 contract: %d required fields, %d cross-field rules | "
               "regional payloads byte-verified: %d of 5 | "
               "physical sensor observations: %d"
               % (v["required_fields"], v["cross_field_rules"],
                  v["regions_byte_verified"], v["physical_observations"])])

    # --- formatting ---
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"].font = Font(italic=True, size=9)
    ws["A3"].font = Font(italic=True, size=9)
    fill = PatternFill("solid", fgColor="DDDDDD")
    for c in ws[5]:
        c.font = Font(bold=True)
        c.fill = fill
    widths = [26, 42, 40, 44, 30, 40, 96]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A6"

    out = os.path.join(DAY6, "GROUP3_DATASET_MAP_STATUS.xlsx")
    wb.save(out)

    print("GROUP 3 - DATASET CONNECTION MAP STATUS")
    print()
    print("  rows tabulated : %d  (5 of Sir's G3 rows + 3 G1 rows where Group 3 is the source)"
          % len(ROWS))
    for r in ROWS:
        print("    %-28s %-34s %s" % (r[0][:28], r[1][:34], r[4]))
    print()
    print("  verified at generation:")
    for k, val in v.items():
        print("    %-24s %s" % (k, val))
    print()
    print("  -> %s" % os.path.basename(out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
